"""
Estimate the what-if ripple engine's links from Castle's history (2015 to today).

    npm run estimate          (exports the hand-set links, then runs this script)

Inputs (local, git-ignored; built by `npm run history -- report`):
    .data/history/sector_monthly.json    per sector-month: ETF total return, deal count, deal value (real $), IPO count
    .data/history/macro_monthly.json     Fed funds, 10-year Treasury, Baa - 10y spread, VIX (monthly averages)
    .data/history/coverage_flags.json    weak years / sector-years from the coverage check
    .data/history/handset_links.json     the hand-set links and their numbers: the priors
    .data/bea/io/*.xlsx                  BEA input-output use tables (AllTablesIO.zip from apps.bea.gov)
Outputs:
    lib/ripple/estimated.json            every link's estimate: value, 90% range, sample size, data weight
    research/link-estimates.md           hand-set vs estimated, link by link, and the biggest surprises

Method, in short (the report says it in plain English):
  * Quarterly panel, 5 sectors. Valuation = sector ETF log total return; deal count, deal value and IPO count enter as
    log changes, log(x + c), with c = 1 for counts and $100M (about one small deal) for values. Weak sector-years from
    the coverage check have their deal figures blanked. Deal and IPO growth are winsorised at the 2.5th / 97.5th
    percentiles (zero quarters and coverage gaps create extreme log changes).
  * C and M: one regression per target feature, pooled across sectors with sector fixed effects (statsmodels OLS),
    on the other features at lags 0-2, the target's own lags 1-2 (controls), and macro changes at lags 0-2. A macro
    effect gets sector-specific coefficients only when an F-test rejects a common effect (p < 0.05).
  * L: half from BEA input-output demand linkages (how much of sector i's output is driven by sector j's final
    demand, Leontief total requirements, 2015-2023 tables), half from the co-movement of the sector ETFs beyond the
    market (partial regression on monthly returns).
  * Shrinkage: every coefficient has a normal prior centred on its hand-set value (0 where there is no hand-set
    link), sd = half its size. The posterior mean is a ridge estimate pulled toward the prior, harder where the data
    is noisy. "Data weight" = 1 - posterior variance / prior variance; links under 0.5 keep the assumption badge.
  * Ranges: 1,000 bootstrap runs resampling whole quarters (all sectors together). Each run re-estimates the
    posterior mean and adds the prior's remaining uncertainty, (1 - weight) x prior sd, so links that are still
    mostly prior keep wide ranges. 90% range = 5th to 95th percentile.
"""

from __future__ import annotations

import json
import math
import sys
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import statsmodels.api as sm

ROOT = Path(__file__).resolve().parent.parent
HIST = ROOT / ".data" / "history"
BEA = ROOT / ".data" / "bea" / "io"
OUT_JSON = ROOT / "lib" / "ripple" / "estimated.json"
OUT_REPORT = ROOT / "research" / "link-estimates.md"

SECTORS = ["Tech", "Healthcare", "Energy", "Financials", "Industrials"]
FEATURES = ["valuation", "dealCount", "dealValue", "ipoCount"]
BASE_MACROS = ["fedFunds", "treasury10y", "baaSpread", "vix"]
# Extra inputs tried in October 2026: the oil price (WTI) and the biotech index (XBI), both in macro_monthly.json. They
# didn't raise the backtest's direction rate, so they're off and the engine has no links for them
# (research/experiments.md lists the hand-set links used). `python research/backtest.py --extra-inputs oil,biotech`.
EXTRA_INPUTS: list[str] = []
MACROS = BASE_MACROS + EXTRA_INPUTS
# Prices enter as % changes (log), the rest as differences. Sector links are controlled for the base inputs only.
MACRO_LOG = {"oil", "biotech"}
# Sector-targeted inputs: estimated only for the sectors a hand-set link names, each against its own prior (a pooled
# coefficient would average oil's Energy effect over five sectors and spread it to all of them).
TARGETED = {"oil", "biotech"}
CONTROL_MACROS = BASE_MACROS
LAGS = [0, 1, 2]
MACRO_COLUMN = {"fedFunds": "fedFundsPct", "treasury10y": "treasury10yPct", "baaSpread": "baaSpreadPct", "vix": "vix", "oil": "oilUsd", "biotech": "biotech"}
MACRO_UNIT = {"fedFunds": 1.0, "treasury10y": 1.0, "baaSpread": 1.0, "vix": 10.0, "oil": 10.0, "biotech": 10.0}  # as MACRO_UNITS in links.ts


def macro_change(frame: pd.DataFrame, m: str) -> pd.Series:
    """One macro input's period-to-period change in its unit (+1 pp; +10 VIX points; +10% for prices)."""
    col = frame[MACRO_COLUMN[m]]
    change = 100 * np.log(col.astype(float)).diff() if m in MACRO_LOG else col.diff()
    return change / MACRO_UNIT[m]

COUNT_CONSTANT = 1.0
VALUE_CONSTANT = 1e8
WINSOR = (2.5, 97.5)
DEAL_VALUE = "total"  # or "capped": estimate deal-value links on deals capped at $5B each (lib/history.ts DEAL_VALUE_CAP_USD)
MACRO_TIMING = "average"  # or "end": macro changes between quarter-end months
# Quarters left out of the estimation (as the quarter explained), e.g. the pandemic shutdown: ["2020Q1", "2020Q2", "2020Q3"].
EXCLUDE_QUARTERS: list[str] = []
# Recency weighting of the C / M equations: a quarter's weight halves every HALF_LIFE years back (None = equal weights).
HALF_LIFE: float | None = None
# Valuation → deal / IPO links may differ by sector (pharma buys biotech when prices fall; energy deals follow oil):
# where an F-test rejects one common coefficient, each sector gets its own, shrunk toward the same hand-set prior.
SECTOR_FEATURES = False
SECTOR_FEATURE_TARGETS = ("dealCount", "dealValue", "ipoCount")
N_BOOT = 1000
SEED = 20261001
ENOUGH_DATA = 0.5  # data weight at which a link drops the "assumption" badge
SEASONAL = False  # quarter-of-year controls in every regression: tested in the backtest, no gain (0.106 vs 0.108), so off


def parse_scale(label) -> dict:
    """Prior width multipliers per link kind: "0.5" (all kinds) or "f0.25/s0.5/m1" (feature / sector / macro links)."""
    text = str(label)
    if "/" not in text:
        v = float(text)
        return {"feature": v, "sector": v, "macro": v}
    parts = dict((p[0], float(p[1:])) for p in text.split("/"))
    return {"feature": parts["f"], "sector": parts["s"], "macro": parts["m"]}


def scale_words(scale: dict) -> str:
    vals = {round(v, 4) for v in scale.values()}
    if len(vals) == 1:
        return f"{scale['feature']:g}"
    return f"{scale['feature']:g} for feature links, {scale['sector']:g} for sector links and {scale['macro']:g} for macro links"


def chosen_prior_scale() -> dict:
    """Prior sd multipliers: the widths that forecast best in the backtest (npm run backtest), else 1 for every kind."""
    p = ROOT / "lib" / "ripple" / "backtest.json"
    if p.exists():
        chosen = json.loads(p.read_text(encoding="utf-8")).get("chosen", {})
        if chosen.get("linkSet") == "estimated" and chosen.get("priorScale"):
            return parse_scale(chosen["priorScale"])
    return parse_scale(1)


PRIOR_SCALE = chosen_prior_scale()
SECTOR_SPECIFIC_P = 0.05
IO_YEARS = list(range(2015, 2024))

LABEL = {"valuation": "Valuation", "dealCount": "Deal count", "dealValue": "Deal value", "ipoCount": "IPO count",
         "fedFunds": "Fed funds", "treasury10y": "10-year yield", "baaSpread": "Baa spread", "vix": "VIX", "oil": "Oil price", "biotech": "Biotech index"}
LAG_TEXT = {0: "same quarter", 1: "1 quarter later", 2: "2 quarters later"}
# "Too weak to matter": under a quarter of the weakest hand-set strength for the link's kind.
TOO_WEAK = {"feature": 0.05, "sector": 0.025, "macro": 0.5}


def prior_sd(kind: str, prior: float) -> float:
    """Half the hand-set size, with a floor; links with no hand-set value get a tight prior around zero."""
    if kind == "feature":
        return max(0.5 * abs(prior), 0.1) if prior else 0.2
    if kind == "own":
        return 0.5
    if kind == "control":
        return 25.0  # missing-data flags: effectively unrestricted
    if kind == "season":
        return 100.0  # quarter-of-year effects (IPOs bunch in Q2, deals dip in Q1): unrestricted
    if kind == "macro":
        return max(0.5 * abs(prior), 1.0)
    if kind == "sector":
        return max(0.5 * abs(prior), 0.05) if prior else 0.1
    raise ValueError(kind)


def log(msg: str) -> None:
    print(msg, flush=True)


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------


def load_json(name: str, fallback=None):
    p = HIST / name
    if not p.exists():
        if fallback is not None:
            return fallback
        sys.exit(f"missing {p}: run `npm run history -- report` first")
    return json.loads(p.read_text(encoding="utf-8"))


def quarter_index(month: str) -> int:
    y, m = int(month[:4]), int(month[5:7])
    return y * 4 + (m - 1) // 3


def quarter_label(qi: int) -> str:
    return f"{qi // 4}Q{qi % 4 + 1}"


def build_panel(deal_value: str | None = None, winsor: tuple | None = None, macro_timing: str | None = None):
    """
    The quarterly panel. `deal_value` "total" or "capped" (each deal at most DEAL_VALUE_CAP_USD, so one megadeal can't
    swing a sector's quarter); `winsor` the percentiles deal and IPO growth are clipped at; `macro_timing` "average"
    (change in quarter averages) or "end" (change from one quarter's last month to the next's).
    """
    deal_value = deal_value or DEAL_VALUE
    winsor = winsor or WINSOR
    macro_timing = macro_timing or MACRO_TIMING
    sm_rows = pd.DataFrame(load_json("sector_monthly.json"))
    if deal_value == "capped":
        if "dealValueCappedUsd" not in sm_rows:
            sys.exit("sector_monthly.json has no capped deal values: run `npm run history -- build`")
        sm_rows["dealValueUsd"] = sm_rows["dealValueCappedUsd"]
    macro = pd.DataFrame(load_json("macro_monthly.json"))
    flags = load_json("coverage_flags.json", fallback={})
    weak = {(f["sector"], f["year"]) for f in flags.get("sectorYears", []) if f.get("weak")}

    sm_rows["qi"] = sm_rows["month"].map(quarter_index)
    complete = sm_rows.groupby("qi")["month"].nunique()
    quarters = sorted(q for q, n in complete.items() if n == 3)
    sm_rows = sm_rows[sm_rows["qi"].isin(quarters)].copy()

    ipo_months = sm_rows[(sm_rows["month"] < "2021-01") & (sm_rows["ipoCount"] > 0)]["month"].nunique()
    ipo_ok = ipo_months >= 12  # same test as ipoHistoryReady() in lib/ripple/baseline.ts

    def agg(g: pd.DataFrame) -> pd.Series:
        r = g["etfReturnPct"]
        return pd.Series({
            "ret": 100 * np.log1p(r / 100).sum() if r.notna().all() and len(r) == 3 else np.nan,
            "deals": g["dealCount"].sum(),
            "value": g["dealValueUsd"].sum(),
            "ipos": g["ipoCount"].sum(),
        })

    q = sm_rows.groupby(["sector", "qi"]).apply(agg, include_groups=False).reset_index()
    q["year"] = (q["qi"] // 4).astype(str)
    q["weak"] = [(s, y) in weak for s, y in zip(q["sector"], q["year"])]
    q.loc[q["weak"], ["deals", "value"]] = np.nan
    if not ipo_ok:
        q["ipos"] = np.nan

    grid = pd.MultiIndex.from_product([SECTORS, quarters], names=["sector", "qi"])
    q = q.set_index(["sector", "qi"]).reindex(grid).reset_index()
    q = q.sort_values(["sector", "qi"])
    g = q.groupby("sector")
    q["valuation"] = q["ret"]
    q["dealCount"] = 100 * (np.log(q["deals"] + COUNT_CONSTANT) - g["deals"].shift(1).pipe(lambda s: np.log(s + COUNT_CONSTANT)))
    q["dealValue"] = 100 * (np.log(q["value"] + VALUE_CONSTANT) - g["value"].shift(1).pipe(lambda s: np.log(s + VALUE_CONSTANT)))
    q["ipoCount"] = 100 * (np.log(q["ipos"] + COUNT_CONSTANT) - g["ipos"].shift(1).pipe(lambda s: np.log(s + COUNT_CONSTANT)))
    clips = {}
    for f in ["dealCount", "dealValue", "ipoCount"]:
        vals = q[f].dropna()
        if len(vals) > 20:
            lo, hi = np.percentile(vals, winsor)
            clips[f] = (lo, hi)
            q[f] = q[f].clip(lo, hi)

    macro["qi"] = macro["month"].map(quarter_index)
    grouped = macro[macro["qi"].isin(quarters)].sort_values("month").groupby("qi")[[MACRO_COLUMN[m] for m in MACROS]]
    mq = grouped.last() if macro_timing == "end" else grouped.mean()
    for m in MACROS:
        mq[m] = macro_change(mq, m)
    q = q.merge(mq[MACROS], left_on="qi", right_index=True, how="left")

    # Lags within each sector (the grid has no gaps).
    g = q.groupby("sector")
    for col in FEATURES + MACROS:
        for lag in LAGS:
            q[f"{col}_l{lag}"] = g[col].shift(lag) if lag else q[col]

    info = {
        "quarters": (quarter_label(quarters[0]), quarter_label(quarters[-1])),
        "nQuarters": len(quarters),
        "weak": sorted(weak, key=lambda w: (w[1], SECTORS.index(w[0]))),
        "ipoOk": ipo_ok,
        "ipoMonths": int(ipo_months),
        "winsor": clips,
        "dealValue": deal_value,
        "firstMonth": sm_rows["month"].min(),
        "lastMonth": sm_rows["month"].max(),
    }
    return q, sm_rows, info


# ---------------------------------------------------------------------------
# Shrinkage estimation
# ---------------------------------------------------------------------------


def within(x: np.ndarray, groups: np.ndarray) -> np.ndarray:
    """Subtract each sector's mean (sector fixed effects)."""
    out = x.astype(float).copy()
    for s in np.unique(groups):
        m = groups == s
        out[m] -= out[m].mean(axis=0)
    return out


def posterior(xd: np.ndarray, yd: np.ndarray, mu: np.ndarray, tau: np.ndarray, sigma2: float):
    precision = xd.T @ xd / sigma2 + np.diag(1 / tau**2)
    cov = np.linalg.inv(precision)
    mean = cov @ (xd.T @ yd / sigma2 + mu / tau**2)
    return mean, cov


def residual_variance(xd: np.ndarray, yd: np.ndarray, n_groups: int) -> float:
    beta, *_ = np.linalg.lstsq(xd, yd, rcond=None)
    resid = yd - xd @ beta
    rank = np.linalg.matrix_rank(xd)
    dof = max(len(yd) - rank - n_groups, 5)
    return float(resid @ resid / dof)


def ols_with_fe(df: pd.DataFrame, y: str, cols: list[str]):
    dummies = pd.get_dummies(df["sector"], dtype=float)
    return sm.OLS(df[y].to_numpy(float), np.column_stack([df[cols].to_numpy(float), dummies.to_numpy()])).fit()


class Regressor:
    def __init__(self, name: str, kind: str, source: str, lag: int, prior: float, sector: str | None = None, scale: float = 1.0):
        self.name, self.kind, self.source, self.lag, self.sector = name, kind, source, lag, sector
        self.prior = prior
        self.tau = prior_sd("feature" if kind == "absent" else kind, prior) * (1.0 if kind in ("control", "season") else scale)


def estimate_equation(panel: pd.DataFrame, target: str, handset: dict, rng: np.random.Generator, n_boot: int = None, prior_scale=None, exclude: list[str] | None = None):
    """
    One target feature: features at lags 0-2, own lags, macro at lags 0-2 (pooled or sector-specific).
    `prior_scale` widens (> 1, trust the data more) or tightens (< 1) the priors, per link kind ("0.5", "f0.25/s0.5/m1"
    or a dict); `n_boot = 0` skips the ranges.
    """
    n_boot = N_BOOT if n_boot is None else n_boot
    scales = PRIOR_SCALE if prior_scale is None else prior_scale if isinstance(prior_scale, dict) else parse_scale(prior_scale)
    scale, macro_scale = scales["feature"], scales["macro"]
    hand_c = {(l["from"], l["to"], l["lag"]): l["value"] for l in handset["features"]}
    hand_m = {}  # (macro, lag, sector) -> hand-set total effect on `target`
    for l in handset["macro"]:
        if l["to"] != target:
            continue
        for s in SECTORS if l["sectors"] == "all" else l["sectors"]:
            hand_m[(l["from"], l["lag"], s)] = hand_m.get((l["from"], l["lag"], s), 0.0) + l["value"]

    available = [f for f in FEATURES if panel[f].notna().sum() > 40]
    if target not in available:
        return None
    feats: list[Regressor] = []
    for g_ in available:
        for lag in LAGS:
            if g_ == target and lag == 0:
                continue
            if g_ == target:
                feats.append(Regressor(f"{g_}_l{lag}", "own", g_, lag, 0.0))
            else:
                prior = hand_c.get((g_, target, lag), 0.0)
                feats.append(Regressor(f"{g_}_l{lag}", "feature" if prior else "absent", g_, lag, prior, scale=scale))
    macro_cols = [f"{m}_l{lag}" for m in MACROS for lag in LAGS]
    # Deal and IPO inputs that are blank (weak sector-years, IPO gaps) count as "no change", with a flag per feature, so
    # the quarter's ETF and macro data still count. The target itself must be present.
    imputable = {"dealCount", "dealValue", "ipoCount"}
    required = [r.name for r in feats if r.source not in imputable] + macro_cols
    df = panel.dropna(subset=[f"{target}_l0"] + required).copy()
    excluded = {quarter_index(f"{q[:4]}-{(int(q[5]) - 1) * 3 + 1:02d}") for q in (EXCLUDE_QUARTERS if exclude is None else exclude)}
    if excluded:
        df = df[~df["qi"].isin(excluded)]
    flags: list[pd.Series] = []
    for g_ in sorted(imputable & set(available)):
        cols = [r.name for r in feats if r.source == g_]
        flag = df[cols].isna().any(axis=1).astype(float)
        if flag.any():
            df[cols] = df[cols].fillna(0.0)
            if not any(flag.equals(f) for f in flags):  # deal count and value go missing together: one flag
                flags.append(flag)
                df[f"missing_{g_}"] = flag
                feats.append(Regressor(f"missing_{g_}", "control", g_, 0, 0.0))
    if SEASONAL:
        # Quarter of the year (Q1 is the base): the calendar pattern shouldn't be read as a link.
        for q in (1, 2, 3):
            df[f"season_q{q + 1}"] = (df["qi"] % 4 == q).astype(float)
            feats.append(Regressor(f"season_q{q + 1}", "season", "season", 0, 0.0))
    for s in SECTORS:
        for m in MACROS:
            for lag in LAGS:
                df[f"{m}_l{lag}_{s}"] = df[f"{m}_l{lag}"] * (df["sector"] == s)

    def macro_regs(specific: set) -> list[Regressor]:
        out = []
        for m in MACROS:
            for lag in LAGS:
                priors = [hand_m.get((m, lag, s), 0.0) for s in SECTORS]
                if m in TARGETED:
                    out += [Regressor(f"{m}_l{lag}_{s}", "macro", m, lag, p, s, scale=macro_scale) for s, p in zip(SECTORS, priors) if p]
                    continue
                if (m, lag) in specific:
                    out += [Regressor(f"{m}_l{lag}_{s}", "macro", m, lag, p, s, scale=macro_scale) for s, p in zip(SECTORS, priors)]
                else:
                    out.append(Regressor(f"{m}_l{lag}", "macro", m, lag, float(np.mean(priors)), scale=macro_scale))
        return out

    # Sector-specific macro effects only where an F-test rejects a common effect.
    y = f"{target}_l0"
    base_cols = [r.name for r in feats] + [r.name for r in macro_regs(set())]
    base = ols_with_fe(df, y, base_cols)
    tests = {}
    for m in MACROS:
        if m in TARGETED:
            continue
        for lag in LAGS:
            cols = [r.name for r in feats] + [r.name for r in macro_regs({(m, lag)})]
            f_stat, p, _ = ols_with_fe(df, y, cols).compare_f_test(base)
            tests[(m, lag)] = float(p)
    specific = {k for k, p in tests.items() if p < SECTOR_SPECIFIC_P}

    # The same for valuation → this target (deal and IPO targets only, when SECTOR_FEATURES is on).
    def feature_regs(split: set) -> list[Regressor]:
        out = []
        for r in feats:
            if r.source == "valuation" and r.lag in split:
                for s in SECTORS:
                    df[f"{r.name}_{s}"] = df[r.name] * (df["sector"] == s)
                    out.append(Regressor(f"{r.name}_{s}", r.kind, r.source, r.lag, r.prior, s, scale=scale))
            else:
                out.append(r)
        return out

    split = set()
    if SECTOR_FEATURES and target in SECTOR_FEATURE_TARGETS and "valuation" in available:
        for lag in LAGS:
            cols = [r.name for r in feature_regs({lag})] + [r.name for r in macro_regs(set())]
            f_stat, p, _ = ols_with_fe(df, y, cols).compare_f_test(base)
            tests[("valuation", lag)] = float(p)
            if p < SECTOR_SPECIFIC_P:
                split.add(lag)
    regs = feature_regs(split) + macro_regs(specific)
    cols = [r.name for r in regs]

    groups = df["sector"].to_numpy()
    x = df[cols].to_numpy(float)
    yv = df[y].to_numpy(float)
    mu = np.array([r.prior for r in regs])
    tau = np.array([r.tau for r in regs])
    final = ols_with_fe(df, y, cols)
    sigma2 = float(final.ssr / final.df_resid)
    xd, yd = within(x, groups), within(yv[:, None], groups)[:, 0]
    if HALF_LIFE:
        # Recent quarters count more: weight halves every HALF_LIFE years back (mean weight 1).
        qi = df["qi"].to_numpy(float)
        w = 0.5 ** ((qi.max() - qi) / (4 * HALF_LIFE))
        sw = np.sqrt(w / w.mean())
        xd, yd = xd * sw[:, None], yd * sw
    mean, cov = posterior(xd, yd, mu, tau, sigma2)
    weight = np.clip(1 - np.diag(cov) / tau**2, 0, 1)

    # Bootstrap: resample whole quarters (every sector's row for that quarter together).
    qis = df["qi"].to_numpy()
    by_q = {qq: np.flatnonzero(qis == qq) for qq in np.unique(qis)}
    keys = list(by_q)
    draws = np.empty((n_boot, len(regs))) if n_boot else mean[None, :].copy()
    for b in range(n_boot):
        idx = np.concatenate([by_q[k] for k in rng.choice(keys, size=len(keys), replace=True)])
        gb = groups[idx]
        xb, yb = within(x[idx], gb), within(yv[idx][:, None], gb)[:, 0]
        s2 = residual_variance(xb, yb, len(np.unique(gb)))
        mb, _ = posterior(xb, yb, mu, tau, s2)
        draws[b] = mb + rng.normal(0, (1 - weight) * tau)

    r2 = 1 - float(((yd - xd @ np.linalg.lstsq(xd, yd, rcond=None)[0]) ** 2).sum() / (yd**2).sum())
    return {
        "target": target, "regs": regs, "mean": mean, "weight": weight, "draws": draws, "n": len(df),
        "sectorsN": df.groupby("sector").size().reindex(SECTORS, fill_value=0).to_dict(),
        "tests": tests, "specific": specific, "sigma2": sigma2, "r2Within": r2, "olsBeta": final.params[: len(cols)],
        "olsSe": final.bse[: len(cols)], "quarters": len(keys),
    }


# ---------------------------------------------------------------------------
# Sector links: BEA input-output + ETF co-movement
# ---------------------------------------------------------------------------

GOVERNMENT = {"GFGD", "GFGN", "GFE", "GSLG", "GSLE", "Used", "Other"}
SECTOR_OF = {
    **{c: "Tech" for c in ["334", "511", "512", "513", "514", "5415"]},
    **{c: "Healthcare" for c in ["621", "622", "623"]},
    **{c: "Energy" for c in ["211", "212", "213", "22", "324", "331", "486"]},
    **{c: "Financials" for c in ["521CI", "523", "524", "525", "HS", "ORE"]},
}
GROUPS = SECTORS + ["Rest"]


def cell(v) -> float:
    if v is None or v == "..." or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def detail_shares() -> dict:
    """Pharma's share of chemicals (325) and medical devices' share of misc. manufacturing (339), 2017 detail table."""
    wb = openpyxl.load_workbook(BEA / "IOUse_Before_Redefinitions_PRO_2017_Detail.xlsx", read_only=True)
    rows = list(wb["2017"].iter_rows(values_only=True))
    desc_row = next(r for r in rows if r and "Total Commodity Output" in [str(c).strip() if c else "" for c in r])
    col = [str(c).strip() if c else "" for c in desc_row].index("Total Commodity Output")
    out = {}
    for code in rows:
        c = str(code[0] or "").strip()
        if c[:3] in ("325", "339") and len(c) == 6:
            out[c] = cell(code[col])
    total = lambda pre: sum(v for k, v in out.items() if k.startswith(pre))
    pharma = sum(out.get(c, 0) for c in ["325411", "325412", "325413", "325414"]) / total("325")
    medical = sum(out.get(c, 0) for c in ["339112", "339113", "339114", "339115", "339116"]) / total("339")
    return {"325": pharma, "339": medical}


def weights(code: str, split: dict) -> dict:
    if code in GOVERNMENT:
        return {"Rest": 1.0}
    if code == "325":
        return {"Healthcare": split["325"], "Energy": 1 - split["325"]}
    if code == "339":
        return {"Healthcare": split["339"], "Industrials": 1 - split["339"]}
    return {SECTOR_OF.get(code, "Industrials"): 1.0}


def io_elasticities(split: dict) -> dict[int, np.ndarray]:
    """Per year: e[i][j] = share of sector i's output driven by sector j's final demand (Leontief, 5 sectors + rest)."""
    wb = openpyxl.load_workbook(BEA / "IOUse_Before_Redefinitions_PRO_1997-2023_Summary.xlsx", read_only=True)
    out = {}
    for year in IO_YEARS:
        rows = list(wb[str(year)].iter_rows(values_only=True))
        code_row = next(i for i, r in enumerate(rows) if r and "111CA" in [str(c).strip() if c else "" for c in r[:4]])
        codes = [str(c).strip() if c else "" for c in rows[code_row]]
        first = codes.index("111CA")
        industries = []
        for c in codes[first:]:
            if not c or c == "None":
                break
            industries.append(c)
        names = [str(c).strip() if c else "" for c in rows[code_row + 1]]
        out_col = names.index("Total Commodity Output")
        z = np.zeros((6, 6))
        x = np.zeros(6)
        start = code_row + 2
        for r in rows[start:]:
            c = str(r[0] or "").strip()
            if not c or c.startswith("V0") or c == "IOCode":
                if str(r[1] or "").strip().startswith("Total Intermediate"):
                    break
                continue
            wr = weights(c, split)
            total = cell(r[out_col])
            for gi, w in wr.items():
                x[GROUPS.index(gi)] += w * total
            for j, ind in enumerate(industries):
                v = cell(r[first + j])
                if not v:
                    continue
                for gi, w in wr.items():
                    for gj, wj in weights(ind, split).items():
                        z[GROUPS.index(gi), GROUPS.index(gj)] += w * wj * v
        f = x - z.sum(axis=1)
        a = z / x[None, :]
        leontief = np.linalg.inv(np.eye(6) - a)
        e = leontief * f[None, :] / x[:, None]
        out[year] = e[:5, :5]
    return out


def market_slopes(rets: np.ndarray, macro: np.ndarray) -> np.ndarray:
    """
    s[i][j] = how much of sector j's own move shows up in sector i: the slope of i's residual on j's residual, after
    taking out the market (the average of the other three sectors, so neither sector is in its own benchmark) and the
    four macro inputs (they have links of their own).
    """
    n = len(SECTORS)
    out = np.zeros((n, n))
    for i in range(n):
        for j in range(n):
            if i == j:
                continue
            others = [k for k in range(n) if k not in (i, j)]
            z = np.column_stack([np.ones(len(rets)), rets[:, others].mean(axis=1), macro])
            pair = rets[:, [i, j]]
            resid = pair - z @ np.linalg.lstsq(z, pair, rcond=None)[0]
            out[i, j] = float(resid[:, 0] @ resid[:, 1] / (resid[:, 1] @ resid[:, 1]))
    return out


_IO_CACHE: dict = {}


def io_tables() -> tuple[dict, dict]:
    """The BEA split shares and per-year IO elasticities, read once per run (the workbooks are slow to open)."""
    if not _IO_CACHE:
        split = detail_shares()
        _IO_CACHE.update(split=split, io=io_elasticities(split))
    return _IO_CACHE["split"], _IO_CACHE["io"]


def estimate_sector_links(sm_rows: pd.DataFrame, handset: dict, rng: np.random.Generator, n_boot: int = None, prior_scale=None, before_month: str | None = None):
    """
    Sector links from IO tables and ETF co-movement. `before_month` ("2019-01") limits both to what was known then
    (ETF months before it, IO tables up to the year before), for the backtest.
    """
    n_boot = N_BOOT if n_boot is None else n_boot
    scales = PRIOR_SCALE if prior_scale is None else prior_scale if isinstance(prior_scale, dict) else parse_scale(prior_scale)
    scale = scales["sector"]
    split, io_all = io_tables()
    years = [y for y in IO_YEARS if before_month is None or y < int(before_month[:4])] or IO_YEARS[:1]
    io = {y: io_all[y] for y in years}
    io_mean = np.mean([io[y] for y in years], axis=0)
    rets = sm_rows.pivot(index="month", columns="sector", values="etfReturnPct")[SECTORS]
    rets = 100 * np.log1p(rets / 100)
    levels = pd.DataFrame(load_json("macro_monthly.json")).set_index("month")
    macro = pd.DataFrame({MACRO_COLUMN[m]: macro_change(levels, m) for m in CONTROL_MACROS})
    rets = rets.join(macro, how="inner").dropna()
    if before_month:
        rets = rets[rets.index < before_month]
    months = rets.index.to_list()
    blocks = {}
    for m in months:
        blocks.setdefault(quarter_index(m), []).append(m)
    keys = list(blocks)
    hand = {(l["from"], l["to"]): l["value"] for l in handset["sectors"]}
    pairs = [(j, i) for j in SECTORS for i in SECTORS if i != j]  # (from, to)

    r, mkt = rets[SECTORS].to_numpy(), rets[[MACRO_COLUMN[m] for m in CONTROL_MACROS]].to_numpy()
    slopes = market_slopes(r, mkt)
    point = {}
    boot = {p: np.empty(n_boot) for p in pairs}
    for (j, i) in pairs:
        e = io_mean[SECTORS.index(i), SECTORS.index(j)]
        b = slopes[SECTORS.index(i), SECTORS.index(j)]
        point[(j, i)] = (0.5 * e + 0.5 * b, e, b)
    pos = {m: k for k, m in enumerate(months)}
    for k in range(n_boot):
        idx = [pos[m] for q in rng.choice(keys, size=len(keys), replace=True) for m in blocks[q]]
        sb = market_slopes(r[idx], mkt[idx])
        year = years[rng.integers(len(years))]
        for (j, i) in pairs:
            boot[(j, i)][k] = 0.5 * io[year][SECTORS.index(i), SECTORS.index(j)] + 0.5 * sb[SECTORS.index(i), SECTORS.index(j)]

    out = {}
    for p in pairs:
        prior = hand.get(p, 0.0)
        tau = prior_sd("sector", prior) * scale
        d, e, b = point[p]
        s2 = float(np.var(boot[p], ddof=1))
        w = tau**2 / (tau**2 + s2)
        mean = (1 - w) * prior + w * d
        draws = (1 - w) * prior + w * boot[p] + rng.normal(0, (1 - w) * tau, n_boot)
        out[p] = {"value": mean, "draws": draws, "weight": w, "prior": prior, "io": e, "market": b, "data": d, "n": len(months)}
    return out, split


# ---------------------------------------------------------------------------
# From coefficients to links
# ---------------------------------------------------------------------------


def summarise(value: float, draws: np.ndarray) -> tuple[float, float]:
    lo, hi = np.percentile(draws, [5, 95])
    return float(lo), float(hi)


def verdict(kind: str, prior: float, value: float, lo: float, hi: float, weight: float) -> str:
    excludes_zero = lo > 0 or hi < 0
    if prior and np.sign(value) != np.sign(prior) and excludes_zero:
        v = "flipped"
    elif abs(value) < TOO_WEAK[kind] and not excludes_zero:
        v = "too weak"
    elif lo <= prior <= hi:
        v = "agreed"
    elif prior and np.sign(value) == np.sign(prior):
        v = "stronger" if abs(value) > abs(prior) else "weaker"
    else:
        v = "leans the other way"
    return v


def rnd(v: float, kind: str) -> float:
    return round(float(v), 2 if kind == "macro" else 3)


def collect(eqs: dict, sector_links: dict, handset: dict, info: dict):
    links = {"features": [], "sectors": [], "macro": [], "macroCompare": []}
    candidates = []
    # Feature links: the coefficient of `from` at `lag` in the `to` equation.
    for l in handset["features"]:
        eq = eqs.get(l["to"])
        idx = next((k for k, r in enumerate(eq["regs"]) if r.kind in ("feature", "absent") and r.source == l["from"] and r.lag == l["lag"] and r.sector is None), None) if eq else None
        tau = prior_sd("feature", l["value"])
        per_sector = [(k, r) for k, r in enumerate(eq["regs"]) if r.kind in ("feature", "absent") and r.source == l["from"] and r.lag == l["lag"] and r.sector] if eq else []
        if per_sector:
            # One estimate per target sector; the link's own value is their average (what the report compares).
            ks = [k for k, _ in per_sector]
            v = float(np.mean(eq["mean"][ks]))
            lo, hi = summarise(v, eq["draws"][:, ks].mean(axis=1))
            links["features"].append({"key": l["key"], "value": rnd(v, "feature"), "range": [rnd(lo, "feature"), rnd(hi, "feature")], "n": eq["n"],
                                      "weight": round(float(np.mean(eq["weight"][ks])), 2), "prior": l["value"],
                                      "bySector": {r.sector: rnd(eq["mean"][k], "feature") for k, r in per_sector}})
            continue
        if idx is None:
            reason = "IPO history not in yet" if "ipoCount" in (l["from"], l["to"]) and not info["ipoOk"] else \
                "cross-sector co-movement: not estimable in a within-sector panel" if l["from"] == l["to"] else "no data"
            draws = l["value"] + np.random.default_rng(SEED).normal(0, tau, N_BOOT)
            lo, hi = summarise(l["value"], draws)
            links["features"].append({"key": l["key"], "value": l["value"], "range": [rnd(lo, "feature"), rnd(hi, "feature")], "n": 0, "weight": 0.0, "prior": l["value"], "note": reason})
            continue
        v, w = float(eq["mean"][idx]), float(eq["weight"][idx])
        lo, hi = summarise(v, eq["draws"][:, idx])
        links["features"].append({"key": l["key"], "value": rnd(v, "feature"), "range": [rnd(lo, "feature"), rnd(hi, "feature")], "n": eq["n"], "weight": round(w, 2), "prior": l["value"],
                                  "ols": rnd(eq["olsBeta"][idx], "feature"), "olsSe": rnd(eq["olsSe"][idx], "feature")})
    for eq in eqs.values():
        for k, r in enumerate(eq["regs"]):
            if r.kind == "absent":
                lo, hi = summarise(eq["mean"][k], eq["draws"][:, k])
                if abs(eq["mean"][k]) >= 0.2 and (lo > 0 or hi < 0) and eq["weight"][k] >= ENOUGH_DATA:
                    candidates.append({"kind": "feature", "label": f"{LABEL[r.source]} → {LABEL[eq['target']]} ({LAG_TEXT[r.lag]})", "value": float(eq["mean"][k]), "range": (lo, hi), "weight": float(eq["weight"][k])})

    # Sector links.
    hand_pairs = {(l["from"], l["to"]): l for l in handset["sectors"]}
    for (j, i), est in sector_links.items():
        lo, hi = summarise(est["value"], est["draws"])
        if (j, i) in hand_pairs:
            links["sectors"].append({"key": hand_pairs[(j, i)]["key"], "value": rnd(est["value"], "sector"), "range": [rnd(lo, "sector"), rnd(hi, "sector")], "n": est["n"],
                                     "weight": round(est["weight"], 2), "prior": est["prior"], "io": rnd(est["io"], "sector"), "market": rnd(est["market"], "sector"),
                                     "ols": rnd(est["data"], "sector")})
        elif abs(est["value"]) >= 0.1 and (lo > 0 or hi < 0):
            candidates.append({"kind": "sector", "label": f"{j} → {i}", "value": est["value"], "range": (lo, hi), "weight": est["weight"], "io": est["io"], "market": est["market"]})

    # Macro links: sector vector per (macro, target, lag), from the chosen specification.
    groups = {}
    for l in handset["macro"]:
        groups.setdefault((l["from"], l["to"], l["lag"]), []).append(l)
    for eq in eqs.values():
        f = eq["target"]
        for m in MACROS:
            for lag in LAGS:
                idx = [k for k, r in enumerate(eq["regs"]) if r.kind == "macro" and r.source == m and r.lag == lag]
                if m in TARGETED:
                    # One link per named sector, with its own estimate and the hand-set link's reason.
                    for k in idx:
                        s = eq["regs"][k].sector
                        based = next((l for l in groups.get((m, f, lag), []) if l["sectors"] == "all" or s in l["sectors"]), None)
                        if based is None:
                            continue
                        lo, hi = np.percentile(eq["draws"][:, k], [5, 95])
                        entry = {"key": f"M:{m}>{f}@{lag}:{s}", "value": rnd(eq["mean"][k], "macro"), "range": [rnd(lo, "macro"), rnd(hi, "macro")], "n": int(eq["sectorsN"][s]),
                                 "weight": round(float(eq["weight"][k]), 2), "prior": eq["regs"][k].prior, "sectorSpecific": True, "p": None, "ols": rnd(eq["olsBeta"][k], "macro")}
                        links["macroCompare"].append({**entry, "key": based["key"]})
                        links["macro"].append({**entry, "from": m, "to": f, "lag": lag, "sectors": [s], "basedOn": based["key"]})
                    continue
                specific = (m, lag) in eq["specific"]
                beta = np.array([eq["mean"][idx[SECTORS.index(s)]] if specific else eq["mean"][idx[0]] for s in SECTORS])
                ols = np.array([eq["olsBeta"][idx[SECTORS.index(s)]] if specific else eq["olsBeta"][idx[0]] for s in SECTORS])
                bdraws = np.column_stack([eq["draws"][:, idx[SECTORS.index(s)]] if specific else eq["draws"][:, idx[0]] for s in SECTORS])
                wvec = np.array([eq["weight"][idx[SECTORS.index(s)]] if specific else eq["weight"][idx[0]] for s in SECTORS])
                hand_links = groups.get((m, f, lag), [])
                if not hand_links:
                    lo, hi = np.percentile(bdraws.mean(axis=1), [5, 95])
                    if abs(beta.mean()) >= 2 and (lo > 0 or hi < 0) and wvec.mean() >= ENOUGH_DATA:
                        candidates.append({"kind": "macro", "label": f"{LABEL[m]} → {LABEL[f]} ({LAG_TEXT[lag]}{', sector-specific' if specific else ''})", "value": float(beta.mean()), "range": (float(lo), float(hi)), "weight": float(wvec.mean())})
                    continue
                # Projection onto the hand-set structure: for the comparison (and the engine when the effect is common).
                cover = np.array([[1.0 if (l["sectors"] == "all" or s in l["sectors"]) else 0.0 for l in hand_links] for s in SECTORS])
                proj = np.linalg.pinv(cover)
                theta, tdraws, tols = proj @ beta, bdraws @ proj.T, proj @ ols
                for k, l in enumerate(hand_links):
                    secs = SECTORS if l["sectors"] == "all" else l["sectors"]
                    lo, hi = np.percentile(tdraws[:, k], [5, 95])
                    w = float(np.mean([wvec[SECTORS.index(s)] for s in secs]))
                    n = int(sum(eq["sectorsN"][s] for s in secs))
                    entry = {"key": l["key"], "value": rnd(theta[k], "macro"), "range": [rnd(lo, "macro"), rnd(hi, "macro")], "n": n, "weight": round(w, 2), "prior": l["value"],
                             "sectorSpecific": specific, "p": round(eq["tests"][(m, lag)], 3), "ols": rnd(tols[k], "macro")}
                    links["macroCompare"].append(entry)
                    if not specific and abs(theta[k]) > 1e-9:
                        links["macro"].append({**entry, "from": m, "to": f, "lag": lag, "sectors": l["sectors"], "basedOn": l["key"]})
                if specific:
                    # The engine gets one link per sector, each with its own estimate; reasons from the most specific hand-set link.
                    for si, s in enumerate(SECTORS):
                        covering = [l for l in hand_links if l["sectors"] == "all" or s in l["sectors"]]
                        if not covering:
                            continue  # no hand-set link (and no reason) for this sector: none is added
                        based = min(covering, key=lambda l: 5 if l["sectors"] == "all" else len(l["sectors"]))
                        lo, hi = np.percentile(bdraws[:, si], [5, 95])
                        prior = sum(l["value"] for l in hand_links if l["sectors"] == "all" or s in l["sectors"])
                        links["macro"].append({"key": f"M:{m}>{f}@{lag}:{s}", "from": m, "to": f, "lag": lag, "sectors": [s], "basedOn": based["key"], "value": rnd(beta[si], "macro"),
                                               "range": [rnd(lo, "macro"), rnd(hi, "macro")], "n": int(eq["sectorsN"][s]), "weight": round(float(wvec[si]), 2), "prior": prior,
                                               "sectorSpecific": True, "p": round(eq["tests"][(m, lag)], 3)})
    # Hand-set macro links whose target equation couldn't be estimated (IPO history missing): prior only.
    done = {e["key"] for e in links["macroCompare"]}
    for l in handset["macro"]:
        if l["key"] in done:
            continue
        tau = prior_sd("macro", l["value"])
        lo, hi = l["value"] - 1.645 * tau, l["value"] + 1.645 * tau
        entry = {"key": l["key"], "value": l["value"], "range": [rnd(lo, "macro"), rnd(hi, "macro")], "n": 0, "weight": 0.0, "prior": l["value"],
                 "note": "IPO history not in yet" if l["to"] == "ipoCount" else "no data"}
        links["macroCompare"].append(entry)
        links["macro"].append({**entry, "from": l["from"], "to": l["to"], "lag": l["lag"], "sectors": l["sectors"], "basedOn": l["key"]})
    return links, candidates


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------


def describe(l: dict, kind: str) -> str:
    if kind == "feature":
        return f"{LABEL[l['from']]} → {LABEL[l['to']]} ({LAG_TEXT[l['lag']]})"
    if kind == "sector":
        return f"{l['from']} → {l['to']}"
    secs = "all sectors" if l["sectors"] == "all" else ", ".join(l["sectors"])
    return f"{LABEL[l['from']]} → {LABEL[l['to']]} ({secs}, {LAG_TEXT[l['lag']]})"


def fmt(v: float, kind: str) -> str:
    d = 1 if kind == "macro" else 2
    s = f"{abs(v):.{d}f}"
    return ("+" if v > 0 else "−" if v < 0 else "") + s


UNITS = {"feature": "% per 1% move", "sector": "share of a move", "macro": "% per unit (1 pp; 10 VIX points)"}


def plain(l: dict, e: dict, kind: str) -> str:
    """One sentence on what the estimate means, in the link's own terms."""
    v, p, (lo, hi) = e["value"], l["value"], e["range"]
    if kind == "feature":
        say = lambda x: f"{abs(10 * x):.1f}% {'higher' if x >= 0 else 'lower'}"
        return (f"A 10% rise in {LABEL[l['from']].lower()} goes with {say(v)} {LABEL[l['to']].lower()} {LAG_TEXT[l['lag']]} "
                f"(90% range {fmt(10 * lo, 'macro')}% to {fmt(10 * hi, 'macro')}%); we had assumed {say(p)}.")
    if kind == "sector":
        return (f"About {abs(v):.0%} of a move in {l['from']} {'reaches' if v >= 0 else 'shows up reversed in'} {l['to']} (90% range {fmt(lo, kind)} to {fmt(hi, kind)}); "
                f"we had assumed {p:.0%}. Input-output links alone say {fmt(e['io'], kind)}, the ETFs' co-movement beyond the market {fmt(e['market'], kind)}.")
    unit = "10 VIX points" if l["from"] == "vix" else "10%" if l["from"] in MACRO_LOG else "1 percentage point"
    secs = "" if l["sectors"] == "all" else f" in {', '.join(l['sectors'])}"
    return (f"A {unit} rise in {LABEL[l['from']]} goes with {fmt(v, kind)}% {LABEL[l['to']].lower()}{secs} {LAG_TEXT[l['lag']]} "
            f"(90% range {fmt(lo, kind)}% to {fmt(hi, kind)}%); we had assumed {fmt(p, kind)}%.")


def write_report(handset: dict, links: dict, candidates: list, eqs: dict, sector_links: dict, split: dict, info: dict) -> list[dict]:
    rows = []
    by_key = {e["key"]: e for e in links["features"] + links["sectors"] + links["macroCompare"]}
    for kind, group in (("feature", handset["features"]), ("sector", handset["sectors"]), ("macro", handset["macro"])):
        for l in group:
            e = by_key[l["key"]]
            lo, hi = e["range"]
            v = verdict(kind, l["value"], e["value"], lo, hi, e["weight"])
            if kind == "macro" and l["sectors"] != "all" and e["value"] == 0 and e["range"] == [0, 0] and not e.get("sectorSpecific"):
                v = "merged"
            rows.append({"kind": kind, "link": l, "est": e, "verdict": v, "mostlyPrior": e["weight"] < ENOUGH_DATA,
                         "surprise": abs(e["value"] - l["value"]) / prior_sd(kind, l["value"]) * e["weight"]})

    estimated = [r for r in rows if not r["mostlyPrior"]]
    count = lambda v, rs=estimated: sum(1 for r in rs if r["verdict"] == v)
    lines = [
        "# Ripple links: hand-set vs estimated",
        "",
        f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} by `research/estimate_links.py` (`npm run estimate`). Do not edit by hand.",
        "",
        "## Summary",
        "",
        f"- Data: {info['nQuarters']} quarters ({info['quarters'][0]} to {info['quarters'][1]}) × 5 sectors, from `sector_monthly` and `macro_monthly`.",
        f"- {len(estimated)} of {len(rows)} links now rest mainly on the data (data weight ≥ {ENOUGH_DATA:.0%}) and drop the \"assumption\" badge; {len(rows) - len(estimated)} are still mostly the hand-set prior and keep it.",
        f"- Among the data-driven links: {count('agreed')} agree with the hand-set value (inside the 90% range), {count('stronger')} are stronger and {count('weaker')} weaker than assumed, "
        f"{count('flipped')} flipped direction, {count('leans the other way')} lean the other way without a clear sign, and {count('too weak')} turned out too weak to matter.",
    ]
    merged = [r for r in rows if r["verdict"] == "merged"]
    if merged:
        lines.append(f"- {len(merged)} hand-set sector-specific macro link{'s' if len(merged) > 1 else ''} merged into the common effect: the data finds no separate sector effect ("
                     + "; ".join(describe(r["link"], "macro") for r in merged) + ").")
    if not info["ipoOk"]:
        lines.append(f"- **IPO history is not in yet** ({info['ipoMonths']} months with IPOs before 2021; 12 needed), so every IPO link is still its hand-set prior. Re-run after `npm run history -- ipos`.")
    if info["weak"]:
        weak_by_year = {}
        for s, y in info["weak"]:
            weak_by_year.setdefault(y, []).append(s)
        lines.append("- Weak sector-years left out of the deal figures: " + "; ".join(f"{y}: {', '.join(ss) if len(ss) < 5 else 'all sectors'}" for y, ss in sorted(weak_by_year.items())) + ".")
    else:
        lines.append("- No weak sector-years flagged by the coverage check.")

    surprises = sorted([r for r in rows if not r["mostlyPrior"]], key=lambda r: (r["verdict"] != "flipped", -r["surprise"]))[:3]
    lines += ["", "## The biggest surprises", "",
              "Ranked by how far the data moved a link from its hand-set value (in prior standard deviations, times its data weight); flips first.", ""]
    for i, r in enumerate(surprises, 1):
        l, e, kind = r["link"], r["est"], r["kind"]
        lines.append(f"{i}. **{describe(l, kind)}** ({r['verdict']}). {plain(l, e, kind)} We had assumed: \"{l['why']}\"")
    lines.append("")

    verdict_text = {"agreed": "agreed", "stronger": "stronger than assumed", "weaker": "weaker than assumed", "flipped": "**flipped**", "too weak": "too weak to matter",
                    "leans the other way": "leans the other way", "merged": "merged into the common effect"}
    for kind, title in (("feature", "Feature links (C)"), ("sector", "Sector links (L)"), ("macro", "Macro links (M)")):
        alone = {"feature": "the OLS coefficient (± standard error) before shrinkage", "sector": "the IO / ETF blend before shrinkage", "macro": "the OLS coefficient before shrinkage"}[kind]
        lines += [f"## {title}", "", f"Units: {UNITS[kind]}. Data alone = {alone}. Data weight = how much of the estimate comes from the data rather than the prior.", ""]
        extra = " | IO | ETF" if kind == "sector" else ""
        lines += [f"| Link | Hand-set | Data alone | Estimated | 90% range | n | Data weight | Verdict{extra} |", "|---|---|---|---|---|---|---|---" + ("|---|---" if kind == "sector" else "") + "|"]
        for r in [r for r in rows if r["kind"] == kind]:
            l, e = r["link"], r["est"]
            v = verdict_text[r["verdict"]] + (" (mostly prior)" if r["mostlyPrior"] else "")
            if e.get("note"):
                v = f"prior only: {e['note']}"
            if kind == "macro" and e.get("sectorSpecific"):
                v += f"; sector-specific (F-test p = {e['p']})"
            elif kind == "macro" and l["sectors"] != "all" and e["value"] == 0 and e["range"] == [0, 0]:
                v = f"merged into the common effect: no separate sector effect in the data (F-test p = {e['p']})"
            data_alone = "—" if e.get("ols") is None else fmt(e["ols"], kind) + (f" ± {abs(e['olsSe']):.2f}" if e.get("olsSe") is not None else "")
            ex = f" | {fmt(e['io'], kind)} | {fmt(e['market'], kind)}" if kind == "sector" else ""
            lines.append(f"| {describe(l, kind)} | {fmt(l['value'], kind)} | {data_alone} | {fmt(e['value'], kind)} | {fmt(e['range'][0], kind)} to {fmt(e['range'][1], kind)} | {e['n']} | {e['weight']:.0%} | {v}{ex} |")
        lines.append("")

    lines += ["## Links the data suggests that the engine doesn't have", "",
              "Coefficients with no hand-set link (prior 0) that came out clearly non-zero with mostly data behind them. Not added to the engine: each needs a reason first.", ""]
    if candidates:
        for c in sorted(candidates, key=lambda c: -abs(c["value"]) / TOO_WEAK[c["kind"]]):
            extra = f" (IO {fmt(c['io'], 'sector')}, ETF {fmt(c['market'], 'sector')})" if c["kind"] == "sector" else ""
            lines.append(f"- {c['label']}: {fmt(c['value'], c['kind'])}, 90% range {fmt(c['range'][0], c['kind'])} to {fmt(c['range'][1], c['kind'])}, data weight {c['weight']:.0%}{extra}")
    else:
        lines.append("- None.")
    specific = [(eq["target"], m, lag, p) for eq in eqs.values() for (m, lag), p in eq["tests"].items() if (m, lag) in eq["specific"]]
    lines += ["", "## Sector-specific macro effects", "",
              "A macro input gets its own coefficient per sector only when an F-test rejects a common effect (p < 0.05). Supported:", ""]
    lines += [f"- {LABEL[m]} → {LABEL[t]} ({LAG_TEXT[lag]}): p = {p:.3f}" for t, m, lag, p in specific] or ["- None: every macro effect is common to all five sectors."]
    lines += ["", "## Regressions", "", "| Target | Sector-quarters | Quarters | Within R² | Residual sd |", "|---|---|---|---|---|"]
    for eq in eqs.values():
        lines.append(f"| {LABEL[eq['target']]} | {eq['n']} | {eq['quarters']} | {eq['r2Within']:.2f} | {math.sqrt(eq['sigma2']):.1f} |")
    lines += ["", "## Method", "",
              "- **Panel**: quarterly, 5 sectors. Valuation = sector ETF log total return. Deal count, deal value and IPO count are log changes of log(x + c), "
              f"c = 1 for counts and $100M for values. Deal and IPO growth winsorised at the {WINSOR[0]}th / {WINSOR[1]}th percentiles. Macro inputs are quarter-average changes (VIX per 10 points). "
              "Deal figures in weak sector-years are left out where they are the thing explained; as inputs they count as \"no change\" with a missing-data flag, so those quarters' ETF and macro data still count.",
              "- **C and M**: one OLS per target feature with sector fixed effects (statsmodels), on the other features at lags 0–2, its own lags 1–2 as controls, and the four macro changes at lags 0–2. "
              "Lag-0 links are same-quarter associations: the data can't tell which way they run.",
              "- **L**: half from BEA input-output tables (share of sector i's output driven by sector j's final demand, Leontief total requirements, 2015–2023, 71 industries mapped to the five sectors; "
              f"pharma is {split['325']:.0%} of chemicals and medical devices {split['339']:.0%} of misc. manufacturing, from the 2017 detail table), half from monthly sector ETF returns "
              "(slope of sector i's own move on sector j's, after taking out the market, measured as the other three sectors' average so neither is in its own benchmark, and the four macro inputs, which have links of their own).",
              "- **Shrinkage**: normal prior on every coefficient centred on its hand-set value (0 where no link exists), sd half its size (floors: 0.1 feature, 0.05 sector, 1 macro; 0.2 / 0.1 / 1 around zero), "
              f"times {scale_words(PRIOR_SCALE)}: the prior widths that forecast best out of sample (`npm run backtest`, research/backtest.md). The posterior mean pulls toward the prior, harder where the data is noisy.",
              f"- **Ranges**: {N_BOOT} bootstrap runs resampling whole quarters (and a random IO year for L); each adds the prior's remaining uncertainty, (1 − data weight) × prior sd. 90% range = 5th–95th percentile.",
              f"- **Badge**: a link drops \"assumption\" when its data weight is at least {ENOUGH_DATA:.0%}.", ""]
    OUT_REPORT.write_text("\n".join(lines), encoding="utf-8", newline="\n")  # LF on Windows too (.gitattributes)
    return rows


def main() -> None:
    rng = np.random.default_rng(SEED)
    handset = load_json("handset_links.json")
    panel, sm_rows, info = build_panel()
    log(f"panel: {info['nQuarters']} quarters {info['quarters'][0]}–{info['quarters'][1]}, IPO history {'in' if info['ipoOk'] else 'missing'}, {len(info['weak'])} weak sector-years")
    eqs = {}
    for f in FEATURES:
        eq = estimate_equation(panel, f, handset, rng)
        if eq is None:
            log(f"  {f}: not enough data, links stay at the prior")
            continue
        eqs[f] = eq
        log(f"  {f}: {eq['n']} sector-quarters, within R² {eq['r2Within']:.2f}, sector-specific macro: {sorted(eq['specific']) or 'none'}")
    sector_links, split = estimate_sector_links(sm_rows, handset, rng)
    log(f"  sector links: {len(sector_links)} pairs (BEA IO {IO_YEARS[0]}–{IO_YEARS[-1]} + ETF returns)")
    links, candidates = collect(eqs, sector_links, handset, info)
    rows = write_report(handset, links, candidates, eqs, sector_links, split, info)
    meta = {
        "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "window": {"from": info["quarters"][0], "to": info["quarters"][1], "quarters": info["nQuarters"]},
        "ipoHistory": info["ipoOk"],
        "enoughData": ENOUGH_DATA,
        "bootstrapRuns": N_BOOT,
        "estimated": sum(1 for r in rows if not r["mostlyPrior"]),
        "total": len(rows),
    }
    out = {"meta": meta, "features": links["features"], "sectors": links["sectors"], "macro": links["macro"]}
    OUT_JSON.write_text(json.dumps(out, indent=1, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
    log(f"wrote {OUT_JSON.relative_to(ROOT)} ({meta['estimated']} of {meta['total']} links data-driven) and {OUT_REPORT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
